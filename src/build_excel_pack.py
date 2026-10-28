"""
Build excel/Regional_Manager_Pack.xlsx: a one-page view a city manager can use without Power BI.
Pick a city in the yellow cell; every number updates through SUMIFS and INDEX/MATCH formulas.

Run after notebook 03 (needs the simulated data and data/clean/city_profile.csv):
    python src/build_excel_pack.py
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule, CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

os.makedirs("excel", exist_ok=True)
FONT = "Arial"
BLUE_INPUT = Font(name=FONT, color="0000FF", bold=True, size=12)
GREEN_LINK = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=16)
NORMAL = Font(name=FONT)
GREY = Font(name=FONT, color="52514E", size=9, italic=True)
HEAD_FILL = PatternFill("solid", fgColor="DCE8F7")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
GBP = '£#,##0;(£#,##0);-'
PCT = '0.0%;(0.0%);-'

# ------------------------------------------------------------------ data (cleaned like SQL script 03)
orders = pd.read_csv("data/simulated/orders.csv", parse_dates=["order_datetime"], dtype={"promo_code": str})
lines = pd.read_csv("data/simulated/order_lines.csv").drop_duplicates("order_line_id")
lines = lines[lines["unit_price"] > 0]
products = pd.read_csv("data/simulated/products.csv")
stores = pd.read_csv("data/simulated/stores.csv")
city = pd.read_csv("data/clean/city_profile.csv")

s = lines.merge(orders, on="order_id").merge(products[["product_id", "category", "unit_cost"]], on="product_id")
s = s.merge(stores[["store_id", "city"]], on="store_id")
s["net_sales"] = s["unit_price"] * s["quantity"] * (1 - s["discount_pct"])
s["cogs"] = s["unit_cost"] * s["quantity"]
s["year"] = s["order_datetime"].dt.year.astype(str)
s["month"] = s["order_datetime"].dt.strftime("%Y-%m")
data = (s.groupby(["month", "year", "city", "channel", "category"])
          .agg(net_sales=("net_sales", "sum"), cogs=("cogs", "sum"), units=("quantity", "sum"))
          .reset_index().round(2))
data["month_no"] = data["month"].str[5:].astype(int)
data["city"] = data["city"].replace({"Online": "Online (website)"})
CATS = sorted(data["category"].unique())
CITIES = list(city["city"]) + ["Online (website)"]
YEARS = ["2024", "2025", "2026"]

wb = Workbook()

# ------------------------------------------------------------------ sheet: Data
ws = wb.active
ws.title = "Data"
headers = ["Month", "Year", "Location", "Channel", "Category", "Net sales (£)", "Cost of goods (£)", "Units", "Month no"]
ws.append(headers)
for row in data.itertuples(index=False):
    ws.append(list(row))
n = ws.max_row
for c in range(1, len(headers) + 1):
    ws.cell(1, c).font = BOLD
    ws.cell(1, c).fill = HEAD_FILL
    ws.column_dimensions[get_column_letter(c)].width = 16
for r in range(2, n + 1):
    ws.cell(r, 6).number_format = GBP
    ws.cell(r, 7).number_format = GBP
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:I{n}"
RNG = {k: f"Data!${col}$2:${col}${n}" for k, col in
       {"year": "B", "loc": "C", "cat": "E", "sales": "F", "cogs": "G", "units": "H"}.items()}

# ------------------------------------------------------------------ sheet: City Facts (real data)
cf = wb.create_sheet("City Facts")
facts = city[["city", "nation", "population", "gdhi_per_head", "median_pay", "retail_gva_gbp_m",
              "retail_gva_is_estimate", "retail_pull_index"]]
cf_headers = ["City", "Nation", "Population, city-region (mid-2025)", "Disposable income per head (£, 2024)",
              "Median annual pay, core city (£, 2025)", "Retail GVA (£m, 2024)", "Retail GVA is an estimate?", "Retail pull index"]
cf.append(cf_headers)
for row in facts.itertuples(index=False):
    cf.append([row.city, row.nation, int(row.population), round(row.gdhi_per_head, 0),
               None if pd.isna(row.median_pay) else row.median_pay, round(row.retail_gva_gbp_m, 1),
               "Yes" if row.retail_gva_is_estimate else "No", round(row.retail_pull_index, 2)])
for c in range(1, len(cf_headers) + 1):
    cf.cell(1, c).font = BOLD
    cf.cell(1, c).fill = HEAD_FILL
    cf.cell(1, c).alignment = Alignment(wrap_text=True, vertical="top")
    cf.column_dimensions[get_column_letter(c)].width = 18
cf.row_dimensions[1].height = 45
for r in range(2, cf.max_row + 1):
    cf.cell(r, 3).number_format = "#,##0"
    for c in (4, 5):
        cf.cell(r, c).number_format = GBP
    cf.cell(r, 6).number_format = "#,##0.0"
    for c in range(1, 9):
        cf.cell(r, c).font = NORMAL
notes = [
    "Sources (all real, official statistics):",
    "Each city is measured as its city-region (e.g. London = Greater London, Manchester = the 10 Greater Manchester boroughs). Council lists: data/clean/city_region_councils.csv.",
    "Population: ONS mid-2025 population estimates (released 1 Oct 2026). Dublin: CSO Population and Migration Estimates, April 2026 (Dublin region).",
    "Disposable income: ONS regional GDHI per head 2024 (released 19 Aug 2026). Dublin: CSO disposable income per person 2023, converted at £0.87 per €1.",
    "Median pay: ONS ASHE 2025 provisional, Table 8.7a (by home location). Belfast uses the Northern Ireland figure; not available for Dublin.",
    "Retail GVA: ONS regional GVA (balanced) by industry, SIC 47, current prices (released 23 Sep 2026). ONS suppresses some councils (all of the Bristol and Cardiff city-regions, most of Glasgow's) and does not cover Dublin, so those are estimated (see notebook 01).",
    "Retail pull index: share of retail GVA ÷ share of spending power. Above 1 = the city-region attracts shoppers from outside.",
]
for i, t in enumerate(notes):
    cf.cell(cf.max_row + 2 if i == 0 else cf.max_row + 1, 1, t).font = GREY if i else BOLD

# ------------------------------------------------------------------ sheet: Manager View
mv = wb.create_sheet("Manager View", 0)
mv.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", [30, 16, 16, 16, 14, 4, 30, 16]):
    mv.column_dimensions[col].width = w

mv["A1"] = "Hearthline · Regional Manager Pack"
mv["A1"].font = TITLE
mv["A2"] = "Simulated sales data (calibrated to real ONS/CSO statistics). Pick a location in the yellow cell."
mv["A2"].font = GREY

mv["A4"] = "Location"
mv["A4"].font = BOLD
mv["B4"] = "Manchester"
mv["B4"].font = BLUE_INPUT
mv["B4"].fill = YELLOW
dv = DataValidation(type="list", formula1=f"\"{','.join(CITIES)}\"", allow_blank=False)
dv.error, dv.errorTitle = "Choose a location from the list", "Invalid location"
mv.add_data_validation(dv)
dv.add("B4")
mv["C4"] = "◀ choose from the drop-down"
mv["C4"].font = GREY

# KPI block by year
mv["A6"] = "Key figures"
mv["A6"].font = BOLD
kpi_head = ["", "2024", "2025", "2026 (Jan–Sep)", "2025 → 2026 like-for-like"]
for i, h in enumerate(kpi_head):
    c = mv.cell(7, 1 + i, h)
    c.font = BOLD
    c.fill = HEAD_FILL
    c.alignment = Alignment(horizontal="center", wrap_text=True)
mv.row_dimensions[7].height = 30
MONTH_NO = f"Data!$I$2:$I${n}"
rows = [
    ("Net sales (£)", lambda y: f'=SUMIFS({RNG["sales"]},{RNG["loc"]},$B$4,{RNG["year"]},"{y}")', GBP),
    ("Gross margin (£)", lambda y: f'=SUMIFS({RNG["sales"]},{RNG["loc"]},$B$4,{RNG["year"]},"{y}")-SUMIFS({RNG["cogs"]},{RNG["loc"]},$B$4,{RNG["year"]},"{y}")', GBP),
    ("Gross margin %", lambda y: None, PCT),
    ("Units sold", lambda y: f'=SUMIFS({RNG["units"]},{RNG["loc"]},$B$4,{RNG["year"]},"{y}")', "#,##0"),
    ("Share of company sales", lambda y: f'=IFERROR(SUMIFS({RNG["sales"]},{RNG["loc"]},$B$4,{RNG["year"]},"{y}")/SUMIFS({RNG["sales"]},{RNG["year"]},"{y}"),0)', PCT),
]
for i, (label, f, fmt) in enumerate(rows):
    r = 8 + i
    mv.cell(r, 1, label).font = NORMAL
    for j, y in enumerate(YEARS):
        cell = mv.cell(r, 2 + j)
        if label == "Gross margin %":
            col = get_column_letter(2 + j)
            cell.value = f"=IFERROR({col}9/{col}8,0)"
        else:
            cell.value = f(y)
        cell.number_format = fmt
        cell.font = NORMAL
        cell.border = THIN
# like-for-like growth on net sales: Jan–Sep 2026 vs Jan–Sep 2025
mv["E8"] = f'=IFERROR(D8/SUMIFS({RNG["sales"]},{RNG["loc"]},$B$4,{RNG["year"]},"2025",{MONTH_NO},"<=9")-1,0)'
mv["E8"].number_format = PCT
mv["E8"].font = BOLD
mv["E13"] = "Like-for-like compares Jan–Sep 2026 with Jan–Sep 2025."
mv["E13"].font = GREY

# category table for 2025 (last full year)
mv["A15"] = "Sales by category, 2025 (last full year)"
mv["A15"].font = BOLD
for i, h in enumerate(["Category", "Net sales (£)", "% of location", "Gross margin %", "vs company margin"]):
    c = mv.cell(16, 1 + i, h)
    c.font = BOLD
    c.fill = HEAD_FILL
    c.alignment = Alignment(horizontal="center", wrap_text=True)
first = 17
for i, cat in enumerate(CATS):
    r = first + i
    mv.cell(r, 1, cat).font = NORMAL
    mv.cell(r, 2, f'=SUMIFS({RNG["sales"]},{RNG["loc"]},$B$4,{RNG["year"]},"2025",{RNG["cat"]},$A{r})').number_format = GBP
    mv.cell(r, 3, f"=IFERROR(B{r}/SUM($B${first}:$B${first + len(CATS) - 1}),0)").number_format = PCT
    mv.cell(r, 4, f'=IFERROR(1-SUMIFS({RNG["cogs"]},{RNG["loc"]},$B$4,{RNG["year"]},"2025",{RNG["cat"]},$A{r})/B{r},0)').number_format = PCT
    mv.cell(r, 5, f'=IFERROR(D{r}-(1-SUMIFS({RNG["cogs"]},{RNG["year"]},"2025",{RNG["cat"]},$A{r})/SUMIFS({RNG["sales"]},{RNG["year"]},"2025",{RNG["cat"]},$A{r})),0)').number_format = '+0.0%;-0.0%;0.0%'
    for c in range(1, 6):
        mv.cell(r, c).border = THIN
        mv.cell(r, c).font = NORMAL
last = first + len(CATS) - 1
tot = last + 1
mv.cell(tot, 1, "Total").font = BOLD
mv.cell(tot, 2, f"=SUM(B{first}:B{last})").number_format = GBP
mv.cell(tot, 2).font = BOLD
mv.conditional_formatting.add(f"B{first}:B{last}", DataBarRule(start_type="min", end_type="max", color="2A78D6"))
mv.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))
mv.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, color="008000")))

# real city facts via INDEX/MATCH
mv["G4"] = "About this city (real official statistics)"
mv["G4"].font = BOLD
facts_rows = [("Population, city-region (mid-2025)", 3, "#,##0"), ("Disposable income per head (£)", 4, GBP),
              ("Median annual pay, core city (£, 2025)", 5, GBP), ("Retail GVA (£m, 2024)", 6, "#,##0.0"),
              ("Retail GVA estimated?", 7, "@"), ("Retail pull index", 8, "0.00")]
for i, (label, col, fmt) in enumerate(facts_rows):
    r = 5 + i
    mv.cell(r, 7, label).font = NORMAL
    letter = get_column_letter(col)
    c = mv.cell(r, 8, f"=IFERROR(INDEX('City Facts'!${letter}$2:${letter}$13,MATCH($B$4,'City Facts'!$A$2:$A$13,0)),\"n/a\")")
    c.number_format = fmt
    c.font = GREEN_LINK
    c.border = THIN
mv["G11"] = "Online (website) has no city facts. Sources on the 'City Facts' sheet."
mv["G11"].font = GREY

# chart of category sales
chart = BarChart()
chart.type = "bar"
chart.style = 10
chart.title = "Net sales by category, 2025"
chart.y_axis.title = "£"
chart.add_data(Reference(mv, min_col=2, min_row=16, max_row=last), titles_from_data=True)
chart.set_categories(Reference(mv, min_col=1, min_row=first, max_row=last))
chart.legend = None
chart.height, chart.width = 8, 15
mv.add_chart(chart, f"G{first - 2}")

# ------------------------------------------------------------------ sheet: Notes
nt = wb.create_sheet("Notes")
for i, t in enumerate([
    "How this workbook works",
    "• Manager View: the yellow cell (B4) is the only input. Every other number is a formula (SUMIFS over the Data sheet, INDEX/MATCH into City Facts).",
    "• Data: monthly net sales, cost of goods and units by location, channel and category, aggregated from the cleaned order lines (same cleaning rules as SQL script 03).",
    "• Font colours: blue = input, black = formula, green = link to another sheet.",
    "• Net sales are after discounts and before refunds. Gross margin = net sales − cost of goods.",
    "• The sales data is SIMULATED for a fictional retailer (Hearthline). City facts are REAL official statistics.",
    "• To analyse further: select the Data sheet and Insert → PivotTable.",
]):
    nt.cell(1 + i, 1, t).font = BOLD if i == 0 else NORMAL
nt.column_dimensions["A"].width = 140

wb.save("excel/Regional_Manager_Pack.xlsx")
print(f"saved, {n - 1:,} data rows")
